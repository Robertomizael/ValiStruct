from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from pathlib import Path
import subprocess
import tempfile
import json
import os
import shutil
import sys
import time
import secrets
import hashlib
import threading
from contextlib import contextmanager

app = Flask(__name__)
APP_STARTED_AT = time.time()
app.config["MAX_CONTENT_LENGTH"] = int(os.environ.get("VALISTRUCT_MAX_UPLOAD_MB", "20")) * 1024 * 1024

# CORS: restringido cuando VALISTRUCT_ALLOWED_ORIGINS está definido.
_allowed_origins_raw = os.environ.get("VALISTRUCT_ALLOWED_ORIGINS", "").strip()
_valistruct_env = os.environ.get("VALISTRUCT_ENV", "development").strip().lower()
if _allowed_origins_raw:
    _allowed_origins = [x.strip() for x in _allowed_origins_raw.split(",") if x.strip()]
    CORS(app, resources={r"/*": {"origins": _allowed_origins}})
elif _valistruct_env in {"production","institutional","beta"}:
    raise RuntimeError("VALISTRUCT_ALLOWED_ORIGINS es obligatorio en modo production/institutional/beta.")
else:
    # Desarrollo local solamente.
    CORS(app)

HERE = os.path.dirname(os.path.abspath(__file__))
R_ENGINE = os.path.join(HERE, "lavaan_engine.R")

def find_rscript():
    return shutil.which("Rscript")

def _compute_auth_guard(roles=None):
    """Protect compute/export routes when institutional auth is enabled."""
    if globals().get("AUTH_ENABLED", False):
        return _require_auth(roles=roles)
    return {"username":"local","role":"admin","auth_disabled":True}, None

R_MAX_N = int(os.environ.get("VALISTRUCT_MONTECARLO_MAX_N","100000"))
R_MAX_REPS = int(os.environ.get("VALISTRUCT_MONTECARLO_MAX_REPS","5000"))


@app.get("/health")
def health():
    rscript = find_rscript()
    if not rscript:
        return jsonify({"ok": False, "error": "Rscript no encontrado"}), 503
    try:
        proc = subprocess.run(
            [rscript, os.path.join(HERE, "health.R")],
            capture_output=True, text=True, timeout=30
        )
        if proc.returncode != 0:
            return jsonify({"ok": False, "error": proc.stderr.strip()}), 503
        return jsonify(json.loads(proc.stdout))
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 503

@app.post("/estimate")
def estimate():
    _user, _auth_err = _compute_auth_guard()
    if _auth_err: return _auth_err
    payload = request.get_json(force=True)
    rscript = find_rscript()
    if not rscript:
        return jsonify({"ok": False, "error": "Rscript no está instalado o no está en PATH"}), 503

    required = ["syntax", "csv_text", "estimator"]
    missing = [k for k in required if not payload.get(k)]
    if missing:
        return jsonify({"ok": False, "error": f"Faltan campos: {', '.join(missing)}"}), 400

    with tempfile.TemporaryDirectory() as td:
        in_json = os.path.join(td, "request.json")
        out_json = os.path.join(td, "response.json")
        with open(in_json, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False)

        proc = subprocess.run(
            [rscript, R_ENGINE, in_json, out_json],
            capture_output=True, text=True, timeout=300
        )
        if proc.returncode != 0:
            return jsonify({
                "ok": False,
                "error": proc.stderr.strip() or proc.stdout.strip() or "Error desconocido del motor R"
            }), 500

        if not os.path.exists(out_json):
            return jsonify({"ok": False, "error": "El motor R no generó respuesta"}), 500

        with open(out_json, "r", encoding="utf-8") as f:
            result = json.load(f)
        return jsonify(result)



@app.post("/advanced")
def advanced():
    _user, _auth_err = _compute_auth_guard()
    if _auth_err: return _auth_err
    payload = request.get_json(force=True)
    rscript = find_rscript()
    if not rscript:
        return jsonify({"ok": False, "error": "Rscript no está instalado o no está en PATH"}), 503
    with tempfile.TemporaryDirectory() as td:
        in_json = os.path.join(td, "request.json")
        out_json = os.path.join(td, "response.json")
        with open(in_json, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False)
        proc = subprocess.run(
            [rscript, os.path.join(HERE, "advanced_engine.R"), in_json, out_json],
            capture_output=True, text=True, timeout=600
        )
        if proc.returncode != 0:
            return jsonify({"ok": False, "error": proc.stderr.strip() or proc.stdout.strip()}), 500
        with open(out_json, "r", encoding="utf-8") as f:
            return jsonify(json.load(f))



@app.post("/report-docx")
def report_docx():
    _user, _auth_err = _compute_auth_guard()
    if _auth_err: return _auth_err
    from io import BytesIO
    from flask import send_file
    try:
        from docx import Document
        from docx.shared import Pt
        from bs4 import BeautifulSoup
    except Exception as e:
        return f"Dependencias DOCX no disponibles: {e}", 500

    payload = request.get_json(force=True)
    title = payload.get("title") or "Reporte ValiStruct"
    html = payload.get("html") or ""
    author = payload.get("author") or "Dr. Roberto Joel Tirado Reyes"
    institution = payload.get("institution") or "Universidad Autónoma de Sinaloa"

    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)

    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(14)

    p = doc.add_paragraph()
    p.add_run(author).bold = True
    doc.add_paragraph(institution)

    soup = BeautifulSoup(html, "html.parser")

    for el in soup.find_all(["h2","h3","p","table"], recursive=True):
        if el.find_parent("table") and el.name != "table":
            continue
        if el.name == "h2":
            doc.add_heading(el.get_text(" ", strip=True), level=1)
        elif el.name == "h3":
            doc.add_heading(el.get_text(" ", strip=True), level=2)
        elif el.name == "p":
            txt = el.get_text(" ", strip=True)
            if txt:
                doc.add_paragraph(txt)
        elif el.name == "table":
            rows = el.find_all("tr")
            if not rows:
                continue
            cols = max(len(r.find_all(["th","td"], recursive=False)) for r in rows)
            table = doc.add_table(rows=1, cols=cols)
            table.style = "Table Grid"
            first = rows[0].find_all(["th","td"], recursive=False)
            for j, cell in enumerate(first):
                table.rows[0].cells[j].text = cell.get_text(" ", strip=True)
            for rr in rows[1:]:
                cells = table.add_row().cells
                vals = rr.find_all(["th","td"], recursive=False)
                for j, cell in enumerate(vals):
                    cells[j].text = cell.get_text(" ", strip=True)

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return send_file(
        bio,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name="ValiStruct_reporte_APA7.docx"
    )



@app.post("/missingness")
def missingness():
    _user, _auth_err = _compute_auth_guard()
    if _auth_err: return _auth_err
    payload = request.get_json(force=True)
    rscript = find_rscript()
    if not rscript:
        return jsonify({"ok": False, "error": "Rscript no está instalado o no está en PATH"}), 503
    with tempfile.TemporaryDirectory() as td:
        in_json = os.path.join(td, "request.json")
        out_json = os.path.join(td, "response.json")
        with open(in_json, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False)
        proc = subprocess.run(
            [rscript, os.path.join(HERE, "missingness_engine.R"), in_json, out_json],
            capture_output=True, text=True, timeout=600
        )
        if proc.returncode != 0:
            return jsonify({"ok": False, "error": proc.stderr.strip() or proc.stdout.strip()}), 500
        with open(out_json, "r", encoding="utf-8") as f:
            return jsonify(json.load(f))


@app.post("/article-docx")
def article_docx():
    _user, _auth_err = _compute_auth_guard()
    if _auth_err: return _auth_err
    from io import BytesIO
    from flask import send_file
    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from bs4 import BeautifulSoup
    except Exception as e:
        return f"Dependencias DOCX no disponibles: {e}", 500

    payload = request.get_json(force=True)
    html = payload.get("html") or ""
    title = payload.get("title") or "Tablas de resultados"
    author = payload.get("author") or "Dr. Roberto Joel Tirado Reyes"
    institution = payload.get("institution") or "Universidad Autónoma de Sinaloa"

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(14)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run(author).bold = True
    p3 = doc.add_paragraph(institution)
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER

    soup = BeautifulSoup(html, "html.parser")
    table_blocks = soup.select(".article-table-block")
    if not table_blocks:
        table_blocks = [soup]

    for block in table_blocks:
        title_el = block.select_one(".apa-table-title")
        if title_el:
            t = doc.add_paragraph()
            rr = t.add_run(title_el.get_text(" ", strip=True))
            rr.bold = True

        table_el = block.find("table")
        if table_el:
            rows = table_el.find_all("tr")
            if rows:
                cols = max(len(r.find_all(["th","td"], recursive=False)) for r in rows)
                table = doc.add_table(rows=1, cols=cols)
                table.style = "Table Grid"
                header = rows[0].find_all(["th","td"], recursive=False)
                for j,c in enumerate(header):
                    table.rows[0].cells[j].text = c.get_text(" ", strip=True)
                    for run in table.rows[0].cells[j].paragraphs[0].runs:
                        run.bold = True
                for rr in rows[1:]:
                    vals = rr.find_all(["th","td"], recursive=False)
                    cells = table.add_row().cells
                    for j,c in enumerate(vals):
                        cells[j].text = c.get_text(" ", strip=True)

        note = block.select_one(".apa-note")
        if note:
            p = doc.add_paragraph()
            r = p.add_run(note.get_text(" ", strip=True))
            r.italic = True
            r.font.size = Pt(9)

        doc.add_paragraph()

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return send_file(
        bio,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name="ValiStruct_tablas_articulo.docx"
    )



@app.post("/export-xlsx")
def export_xlsx():
    _user, _auth_err = _compute_auth_guard()
    if _auth_err: return _auth_err
    from io import BytesIO
    from flask import send_file
    try:
        import xlsxwriter
    except Exception as e:
        return f"Dependencia XLSX no disponible: {e}", 500

    payload = request.get_json(force=True)
    bio = BytesIO()
    wb = xlsxwriter.Workbook(bio, {"in_memory": True})

    fmt_title = wb.add_format({"bold": True, "font_size": 16, "font_color": "#7C1F2A"})
    fmt_head = wb.add_format({"bold": True, "bg_color": "#7C1F2A", "font_color": "white", "border": 1})
    fmt_sub = wb.add_format({"bold": True, "bg_color": "#F3E8EA", "border": 1})
    fmt_cell = wb.add_format({"border": 1})
    fmt_num = wb.add_format({"border": 1, "num_format": "0.000"})
    fmt_wrap = wb.add_format({"border": 1, "text_wrap": True})

    def sheet(name):
        ws = wb.add_worksheet(name[:31])
        ws.freeze_panes(1, 0)
        ws.set_column(0, 0, 24)
        ws.set_column(1, 8, 16)
        return ws

    # Summary
    ws = sheet("Resumen")
    ws.write("A1", "ValiStruct · Resultados consolidados", fmt_title)
    ws.write("A3", "Autor", fmt_sub); ws.write("B3", payload.get("author",""))
    ws.write("A4", "Institución", fmt_sub); ws.write("B4", payload.get("institution",""))
    quality = payload.get("quality") or {}
    ws.write("A6", "Cobertura del proceso", fmt_sub); ws.write("B6", quality.get("coverage",""))
    stages = quality.get("stages") or []
    if stages:
        ws.write_row("A8", ["Etapa","Disponible","Revisión","Detalle"], fmt_head)
        r=8
        for s in stages:
            ws.write_row(r,0,[s.get("name"),"Sí" if s.get("done") else "No","Sí" if s.get("review") else "No",s.get("detail")],fmt_cell)
            r+=1
        ws.set_column(3,3,55)

    # Diagnostics
    diag = payload.get("diagnostics") or {}
    if diag.get("rows"):
        ws=sheet("Diagnostico")
        ws.write_row("A1",["Variable","N válido","Faltantes %","Media","DE","Asimetría","Curtosis","Atípicos","Orientación"],fmt_head)
        r=1
        for x in diag["rows"]:
            vals=[x.get("name"),x.get("n"),x.get("missing"),x.get("mean"),x.get("sd"),x.get("skew"),x.get("kurt"),x.get("outliers"),x.get("status")]
            for c,v in enumerate(vals):
                ws.write(r,c,v,fmt_num if c in [2,3,4,5,6] and isinstance(v,(int,float)) else fmt_cell)
            r+=1
        ws.set_column(8,8,48)

    # Reliability
    rel = payload.get("reliability") or {}
    if rel:
        ws=sheet("Confiabilidad")
        ws.write("A1","Alfa",fmt_sub); ws.write("B1",rel.get("alpha"),fmt_num)
        ws.write("A2","Alfa estandarizado",fmt_sub); ws.write("B2",rel.get("alphaStd"),fmt_num)
        ws.write("A3","Omega preliminar",fmt_sub); ws.write("B3",rel.get("omegaApprox"),fmt_num)
        if rel.get("itemRows"):
            ws.write_row("A5",["Ítem","Media","DE","r ítem-total","α si elimina","Orientación"],fmt_head)
            r=5
            for x in rel["itemRows"]:
                ws.write_row(r,0,[x.get("name"),x.get("mean"),x.get("sd"),x.get("rit"),x.get("aDel"),x.get("flag")],fmt_cell)
                r+=1

    # EFA
    efa = payload.get("efa") or {}
    if efa:
        ws=sheet("AFE")
        kmo=(efa.get("kmo") or {}).get("overall")
        bart=efa.get("bart") or {}
        ws.write("A1","KMO",fmt_sub); ws.write("B1",kmo,fmt_num)
        ws.write("A2","Bartlett χ²",fmt_sub); ws.write("B2",bart.get("chi2"),fmt_num)
        ws.write("A3","Bartlett gl",fmt_sub); ws.write("B3",bart.get("df"),fmt_cell)
        ws.write("A4","Bartlett p",fmt_sub); ws.write("B4",bart.get("p"),fmt_num)
        ws.write("A5","Factores análisis paralelo",fmt_sub); ws.write("B5",efa.get("retainedPA"),fmt_cell)
        loads=efa.get("loadings") or []; names=efa.get("itemNames") or []
        if loads:
            nf=len(loads[0])
            headers=["Ítem"]+[f"F{i+1}" for i in range(nf)]+["h²","Orientación"]
            ws.write_row(7,0,headers,fmt_head)
            comm=efa.get("communalities") or []
            diagrows=efa.get("itemDiag") or []
            for i,row in enumerate(loads):
                vals=[names[i] if i<len(names) else f"I{i+1}"]+row+[comm[i] if i<len(comm) else None,diagrows[i].get("status") if i<len(diagrows) else ""]
                ws.write_row(8+i,0,vals,fmt_cell)

    # Motor Pro
    pro=payload.get("motorPro") or {}
    if pro:
        ws=sheet("Motor Pro")
        fit=pro.get("fit") or {}
        ws.write_row("A1",["Índice","Valor"],fmt_head)
        r=1
        for k in ["chisq","df","pvalue","cfi","tli","rmsea","rmsea.ci.lower","rmsea.ci.upper","srmr","aic","bic"]:
            if k in fit:
                ws.write(r,0,k,fmt_cell); ws.write(r,1,fit.get(k),fmt_num); r+=1
        pars=pro.get("parameters") or []
        if pars:
            r+=1
            ws.write_row(r,0,["lhs","op","rhs","Est.","EE","z","p","Std.all"],fmt_head); r+=1
            for p in pars:
                ws.write_row(r,0,[p.get("lhs"),p.get("op"),p.get("rhs"),p.get("est"),p.get("se"),p.get("z"),p.get("pvalue"),p.get("std_all")],fmt_cell);r+=1

    # History
    hist=payload.get("history") or []
    if hist:
        ws=sheet("Historial")
        ws.write_row("A1",["Fecha","Módulo","Acción","Metadatos"],fmt_head)
        for i,h in enumerate(hist, start=1):
            ws.write_row(i,0,[h.get("timestamp"),h.get("module"),h.get("action"),json.dumps(h.get("meta") or {},ensure_ascii=False)],fmt_wrap)
        ws.set_column(3,3,60)

    wb.close()
    bio.seek(0)
    return send_file(
        bio,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="ValiStruct_resultados_consolidados.xlsx"
    )



@app.post("/sem-montecarlo")
def sem_montecarlo():
    _user, _auth_err = _compute_auth_guard()
    if _auth_err: return _auth_err
    payload = request.get_json(force=True)
    try:
        n = int(payload.get("n",300))
        reps = int(payload.get("reps",500))
    except (TypeError, ValueError):
        return jsonify({"ok":False,"error":"n y reps deben ser enteros"}),400
    if n < 20 or n > R_MAX_N:
        return jsonify({"ok":False,"error":f"n fuera de rango (20–{R_MAX_N})"}),400
    if reps < 10 or reps > R_MAX_REPS:
        return jsonify({"ok":False,"error":f"reps fuera de rango (10–{R_MAX_REPS})"}),400
    rscript = find_rscript()
    if not rscript:
        return jsonify({"ok": False, "error": "Rscript no está instalado o no está en PATH"}), 503
    with tempfile.TemporaryDirectory() as td:
        in_json = os.path.join(td, "request.json")
        out_json = os.path.join(td, "response.json")
        with open(in_json, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False)
        proc = subprocess.run(
            [rscript, os.path.join(HERE, "sem_montecarlo.R"), in_json, out_json],
            capture_output=True, text=True, timeout=1200
        )
        if proc.returncode != 0:
            return jsonify({"ok": False, "error": proc.stderr.strip() or proc.stdout.strip()}), 500
        with open(out_json, "r", encoding="utf-8") as f:
            return jsonify(json.load(f))

@app.post("/model-check")
def model_check():
    _user, _auth_err = _compute_auth_guard()
    if _auth_err: return _auth_err
    payload = request.get_json(force=True)
    rscript = find_rscript()
    if not rscript:
        return jsonify({"ok": False, "error": "Rscript no está instalado o no está en PATH"}), 503
    with tempfile.TemporaryDirectory() as td:
        in_json = os.path.join(td, "request.json")
        out_json = os.path.join(td, "response.json")
        with open(in_json, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False)
        proc = subprocess.run(
            [rscript, os.path.join(HERE, "model_check.R"), in_json, out_json],
            capture_output=True, text=True, timeout=120
        )
        if proc.returncode != 0:
            return jsonify({"ok": False, "error": proc.stderr.strip() or proc.stdout.strip()}), 500
        with open(out_json, "r", encoding="utf-8") as f:
            return jsonify(json.load(f))



@app.post("/xlsx-info")
def xlsx_info():
    _user, _auth_err = _compute_auth_guard()
    if _auth_err: return _auth_err
    try:
        import openpyxl
        from io import BytesIO
    except Exception as e:
        return jsonify({"ok": False, "error": f"openpyxl no disponible: {e}"}), 500
    if "file" not in request.files:
        return jsonify({"ok": False, "error": "No se recibió archivo XLSX"}), 400
    f = request.files["file"]
    try:
        wb = openpyxl.load_workbook(BytesIO(f.read()), read_only=True, data_only=True)
        return jsonify({"ok": True, "sheets": wb.sheetnames})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 400

@app.post("/xlsx-to-csv")
def xlsx_to_csv():
    _user, _auth_err = _compute_auth_guard()
    if _auth_err: return _auth_err
    try:
        import openpyxl, csv
        from io import BytesIO, StringIO
    except Exception as e:
        return jsonify({"ok": False, "error": f"Dependencias XLSX no disponibles: {e}"}), 500
    if "file" not in request.files:
        return jsonify({"ok": False, "error": "No se recibió archivo XLSX"}), 400
    f = request.files["file"]
    sheet = request.form.get("sheet","")
    try:
        wb = openpyxl.load_workbook(BytesIO(f.read()), read_only=True, data_only=True)
        ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[0]]
        out = StringIO()
        writer = csv.writer(out)
        for row in ws.iter_rows(values_only=True):
            writer.writerow(["" if v is None else v for v in row])
        return jsonify({"ok": True, "sheet": ws.title, "csv_text": out.getvalue()})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 400



@app.after_request
def security_headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "SAMEORIGIN"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Cache-Control"] = "no-store"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    response.headers["X-Permitted-Cross-Domain-Policies"] = "none"
    return response



@app.post("/legacy-to-csv")
def legacy_to_csv():
    _user, _auth_err = _compute_auth_guard()
    if _auth_err: return _auth_err
    try:
        import pyreadstat
        import pandas as pd
        from io import BytesIO, StringIO
    except Exception as e:
        return jsonify({"ok": False, "error": f"Dependencias SAV/DTA no disponibles: {e}"}), 500

    if "file" not in request.files:
        return jsonify({"ok": False, "error": "No se recibió archivo"}), 400

    f = request.files["file"]
    name = (f.filename or "").lower()
    suffix = ".sav" if name.endswith(".sav") else ".dta" if name.endswith(".dta") else None
    if suffix is None:
        return jsonify({"ok": False, "error": "Formato no soportado. Use .sav o .dta"}), 400

    import tempfile, os
    tmp = None
    try:
        fd, tmp = tempfile.mkstemp(suffix=suffix)
        os.close(fd)
        f.save(tmp)
        if suffix == ".sav":
            df, meta = pyreadstat.read_sav(tmp, apply_value_formats=False)
            fmt = "SPSS SAV"
        else:
            df, meta = pyreadstat.read_dta(tmp, apply_value_formats=False)
            fmt = "Stata DTA"

        out = StringIO()
        df.to_csv(out, index=False)
        labels = {}
        try:
            for col, lab in zip(meta.column_names, meta.column_labels):
                if lab:
                    labels[col] = lab
        except Exception:
            pass

        return jsonify({
            "ok": True,
            "csv_text": out.getvalue(),
            "rows": int(df.shape[0]),
            "columns": int(df.shape[1]),
            "meta": {"format": fmt, "labels": labels}
        })
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 400
    finally:
        if tmp and os.path.exists(tmp):
            try: os.remove(tmp)
            except Exception: pass



# -----------------------------
# ValiStruct 2.3 optional institutional mode
# -----------------------------
AUTH_ENABLED = os.environ.get("VALISTRUCT_AUTH_ENABLED", "false").lower() == "true"
PROJECT_LIBRARY_ENABLED = os.environ.get("VALISTRUCT_PROJECT_LIBRARY_ENABLED", "false").lower() == "true"
PROJECT_DIR = Path(os.environ.get("VALISTRUCT_PROJECT_DIR", os.path.join(HERE, "institution_projects")))
PROJECT_DIR.mkdir(parents=True, exist_ok=True)
ACTIVE_TOKENS = {}
LOGIN_FAILURES = {}
AUTH_TOKEN_TTL_SECONDS = int(os.environ.get("VALISTRUCT_AUTH_TOKEN_TTL_SECONDS","28800"))
AUTH_MAX_FAILURES = int(os.environ.get("VALISTRUCT_AUTH_MAX_FAILURES","5"))
AUTH_FAILURE_WINDOW = int(os.environ.get("VALISTRUCT_AUTH_FAILURE_WINDOW_SECONDS","900"))
AUTH_PBKDF2_ITERATIONS = int(os.environ.get("VALISTRUCT_PBKDF2_ITERATIONS","310000"))

def _load_users():
    """
    Recommended:
      {"username":{"password_pbkdf2_sha256":"iterations$salt_hex$digest_hex","role":"researcher"}}
    Legacy password_sha256 is accepted temporarily for migration only.
    """
    raw = os.environ.get("VALISTRUCT_USERS_JSON", "{}")
    try:
        obj = json.loads(raw)
        return obj if isinstance(obj, dict) else {}
    except Exception:
        return {}

def _verify_password(password, entry):
    stored = str(entry.get("password_pbkdf2_sha256",""))
    if stored:
        try:
            it_s, salt_hex, digest_hex = stored.split("$",2)
            it = int(it_s)
            derived = hashlib.pbkdf2_hmac("sha256",password.encode("utf-8"),bytes.fromhex(salt_hex),it)
            return secrets.compare_digest(derived.hex(),digest_hex)
        except Exception:
            return False
    legacy = str(entry.get("password_sha256",""))
    if legacy:
        digest = hashlib.sha256(password.encode("utf-8")).hexdigest()
        return secrets.compare_digest(digest,legacy)
    return False

def _auth_user_from_request():
    if not AUTH_ENABLED:
        return {"username":"local","role":"admin","auth_disabled":True}
    auth = request.headers.get("Authorization","")
    if not auth.startswith("Bearer "):
        return None
    token = auth.split(" ",1)[1].strip()
    item = ACTIVE_TOKENS.get(token)
    if not item:
        return None
    now = time.time()
    if float(item.get("expires_at",0)) <= now:
        ACTIVE_TOKENS.pop(token,None)
        return None
    return {k:v for k,v in item.items() if k!="expires_at"}

def _require_auth(roles=None):
    user = _auth_user_from_request()
    if user is None:
        return None,(jsonify({"ok":False,"error":"Autenticación requerida"}),401)
    if roles and user.get("role") not in roles:
        return None,(jsonify({"ok":False,"error":"Permisos insuficientes"}),403)
    return user,None

@app.post("/auth/login")
def auth_login():
    if not AUTH_ENABLED:
        return jsonify({"ok":False,"error":"La autenticación institucional está desactivada en el servidor"}),403
    body = request.get_json(force=True)
    username = str(body.get("username","")).strip()
    password = str(body.get("password",""))
    now = time.time()
    key = f"{request.remote_addr or 'unknown'}:{username}"
    recent = [t for t in LOGIN_FAILURES.get(key,[]) if now-t < AUTH_FAILURE_WINDOW]
    LOGIN_FAILURES[key] = recent
    if len(recent) >= AUTH_MAX_FAILURES:
        return jsonify({"ok":False,"error":"Demasiados intentos fallidos. Intente más tarde."}),429

    users = _load_users()
    entry = users.get(username)
    if not entry or not _verify_password(password,entry):
        LOGIN_FAILURES[key] = recent + [now]
        return jsonify({"ok":False,"error":"Credenciales inválidas"}),401

    LOGIN_FAILURES.pop(key,None)
    role = str(entry.get("role","researcher"))
    token = secrets.token_urlsafe(32)
    ACTIVE_TOKENS[token] = {
        "username":username,"role":role,
        "expires_at":now + AUTH_TOKEN_TTL_SECONDS
    }
    return jsonify({
        "ok":True,"token":token,
        "user":{"username":username,"role":role},
        "expires_in":AUTH_TOKEN_TTL_SECONDS,
        "legacy_password_hash":bool(entry.get("password_sha256"))
    })

@app.post("/auth/logout")
def auth_logout():
    auth = request.headers.get("Authorization", "")
    if auth.startswith("Bearer "):
        ACTIVE_TOKENS.pop(auth.split(" ",1)[1].strip(), None)
    return jsonify({"ok": True})

@app.get("/auth/status")
def auth_status():
    user = _auth_user_from_request()
    return jsonify({"ok": True, "authenticated": user is not None, "user": user})

def _project_path(project_id):
    safe = "".join(c for c in str(project_id) if c.isalnum() or c in "-_")[:80]
    return PROJECT_DIR / f"{safe}.json"

def _read_project_meta(path):
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return {
            "id": data.get("_institution",{}).get("id", path.stem),
            "name": data.get("name") or "Proyecto ValiStruct",
            "owner": data.get("_institution",{}).get("owner",""),
            "updated_at": data.get("_institution",{}).get("updated_at","")
        }
    except Exception:
        return None

@app.get("/projects")
def list_projects_api():
    if not PROJECT_LIBRARY_ENABLED:
        return jsonify({"ok": False, "error": "Biblioteca institucional desactivada"}), 403
    user, err = _require_auth()
    if err: return err
    projects = []
    for path in PROJECT_DIR.glob("*.json"):
        meta = _read_project_meta(path)
        if not meta: continue
        try:
            project_obj = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            project_obj = {}
        if user["role"] == "admin" or meta["owner"] == user["username"] or _can_read_project(user, project_obj):
            projects.append(meta)
    projects.sort(key=lambda x: x.get("updated_at",""), reverse=True)
    return jsonify({"ok": True, "role": user["role"], "projects": projects})

@app.post("/projects")
def create_project_api():
    if not PROJECT_LIBRARY_ENABLED:
        return jsonify({"ok": False, "error": "Biblioteca institucional desactivada"}), 403
    user, err = _require_auth(roles={"researcher","teacher","admin"})
    if err: return err
    body = request.get_json(force=True)
    project = body.get("project") or {}
    pid = secrets.token_hex(8)
    now = __import__("datetime").datetime.utcnow().isoformat(timespec="seconds")+"Z"
    project["_institution"] = {"id": pid, "owner": user["username"], "updated_at": now, "revision": 1}
    _project_path(pid).write_text(json.dumps(project,ensure_ascii=False,indent=2),encoding="utf-8")
    _log_activity(pid,user["username"],"Proyecto creado","")
    return jsonify({"ok": True, "id": pid, "revision": 1})

@app.get("/projects/<project_id>")
def get_project_api(project_id):
    if not PROJECT_LIBRARY_ENABLED:
        return jsonify({"ok": False, "error": "Biblioteca institucional desactivada"}), 403
    user, err = _require_auth()
    if err: return err
    path = _project_path(project_id)
    if not path.exists():
        return jsonify({"ok": False, "error": "Proyecto no encontrado"}), 404
    project = json.loads(path.read_text(encoding="utf-8"))
    if not _can_read_project(user, project):
        return jsonify({"ok": False, "error": "Permisos insuficientes"}), 403
    return jsonify({"ok": True, "project": project})

@app.put("/projects/<project_id>")
def update_project_api(project_id):
    if not PROJECT_LIBRARY_ENABLED:
        return jsonify({"ok":False,"error":"Biblioteca institucional desactivada"}),403
    user,err=_require_auth(roles={"researcher","teacher","admin"})
    if err:return err
    path=_project_path(project_id)
    lock_path=Path(str(path)+".lock")
    try:
        with _file_lock(lock_path):
            if not path.exists():
                return jsonify({"ok":False,"error":"Proyecto no encontrado"}),404
            old=json.loads(path.read_text(encoding="utf-8"))
            owner=old.get("_institution",{}).get("owner")
            if not _can_edit_project(user,old):
                return jsonify({"ok":False,"error":"Permiso de edición requerido"}),403
            body=request.get_json(force=True)
            expected_revision=body.get("expected_revision")
            current_revision=int(old.get("_institution",{}).get("revision",1))
            if expected_revision is not None:
                try:
                    expected_revision=int(expected_revision)
                except (TypeError,ValueError):
                    return jsonify({"ok":False,"error":"expected_revision debe ser entero"}),400
                if expected_revision != current_revision:
                    return jsonify({"ok":False,"error":"Conflicto de edición","conflict":True,"current_revision":current_revision}),409
            project=body.get("project") or {}
            now=__import__("datetime").datetime.utcnow().isoformat(timespec="seconds")+"Z"
            next_revision=current_revision+1
            project["_institution"]={"id":project_id,"owner":owner,"updated_at":now,"revision":next_revision}
            _save_version(project_id,old,user["username"],"Versión automática antes de actualización")
            _atomic_write_text(path,json.dumps(project,ensure_ascii=False,indent=2))
        _log_activity(project_id,user["username"],"Proyecto actualizado",f"revision {next_revision}")
        return jsonify({"ok":True,"revision":next_revision})
    except RuntimeError as e:
        return jsonify({"ok":False,"error":str(e)}),500

@app.delete("/projects/<project_id>")
def delete_project_api(project_id):
    if not PROJECT_LIBRARY_ENABLED:
        return jsonify({"ok": False, "error": "Biblioteca institucional desactivada"}), 403
    user, err = _require_auth(roles={"researcher","teacher","admin"})
    if err: return err
    path = _project_path(project_id)
    if not path.exists():
        return jsonify({"ok": False, "error": "Proyecto no encontrado"}), 404
    old = json.loads(path.read_text(encoding="utf-8"))
    owner = old.get("_institution",{}).get("owner")
    if user["role"] != "admin" and owner != user["username"]:
        return jsonify({"ok": False, "error": "Permisos insuficientes"}), 403
    path.unlink()
    return jsonify({"ok": True})

@app.get("/self-test")
def backend_self_test():
    tests = []
    def add(name,status,detail=""):
        tests.append({"name":name,"status":status,"detail":detail})
    add("Flask API","pass","Backend operativo.")
    add("Rscript","pass" if find_rscript() else "warn","Rscript disponible." if find_rscript() else "Rscript no encontrado.")
    try:
        import openpyxl
        add("XLSX / openpyxl","pass",openpyxl.__version__)
    except Exception as e:
        add("XLSX / openpyxl","fail",str(e))
    try:
        import pyreadstat
        add("SAV/DTA / pyreadstat","pass",getattr(pyreadstat,"__version__","disponible"))
    except Exception as e:
        add("SAV/DTA / pyreadstat","fail",str(e))
    try:
        import xlsxwriter
        add("XLSX export","pass",getattr(xlsxwriter,"__version__","disponible"))
    except Exception as e:
        add("XLSX export","fail",str(e))
    try:
        from docx import Document
        add("DOCX export","pass","python-docx disponible.")
    except Exception as e:
        add("DOCX export","fail",str(e))
    add("Auth mode","pass" if AUTH_ENABLED else "warn","Activado." if AUTH_ENABLED else "Desactivado (valor predeterminado).")
    add("Institution project library","pass" if PROJECT_LIBRARY_ENABLED else "warn","Activada." if PROJECT_LIBRARY_ENABLED else "Desactivada (valor predeterminado).")
    return jsonify({"ok": True, "tests": tests})



# -----------------------------
# ValiStruct 2.4 collaboration / versions
# -----------------------------
VERSION_DIR = PROJECT_DIR / "_versions"
PERMISSION_FILE = PROJECT_DIR / "_permissions.json"
VERSION_DIR.mkdir(parents=True, exist_ok=True)

def _load_permissions():
    try:
        return json.loads(PERMISSION_FILE.read_text(encoding="utf-8")) if PERMISSION_FILE.exists() else {}
    except Exception:
        return {}

def _save_permissions(obj):
    PERMISSION_FILE.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding="utf-8")

def _can_read_project(user, project):
    owner = project.get("_institution",{}).get("owner")
    if user.get("role") == "admin" or owner == user.get("username"):
        return True
    pid = project.get("_institution",{}).get("id")
    perms = _load_permissions().get(pid,{})
    return perms.get(user.get("username")) in ("viewer","editor")

def _can_edit_project(user, project):
    owner = project.get("_institution",{}).get("owner")
    if user.get("role") == "admin" or owner == user.get("username"):
        return True
    pid = project.get("_institution",{}).get("id")
    perms = _load_permissions().get(pid,{})
    return perms.get(user.get("username")) == "editor"

def _version_project_dir(project_id):
    d = VERSION_DIR / str(project_id)
    d.mkdir(parents=True, exist_ok=True)
    return d

def _save_version(project_id, project, username, comment=""):
    d = _version_project_dir(project_id)
    existing = []
    for p in d.glob("*.json"):
        try: existing.append(int(p.stem))
        except Exception: pass
    version = max(existing, default=0) + 1
    payload = {
        "version": version,
        "created_at": __import__("datetime").datetime.utcnow().isoformat(timespec="seconds")+"Z",
        "created_by": username,
        "comment": comment,
        "project": project
    }
    (d/f"{version}.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
    return version

@app.get("/projects/<project_id>/permissions")
def get_project_permissions(project_id):
    user, err = _require_auth()
    if err: return err
    path = _project_path(project_id)
    if not path.exists():
        return jsonify({"ok": False, "error": "Proyecto no encontrado"}), 404
    project = json.loads(path.read_text(encoding="utf-8"))
    owner = project.get("_institution",{}).get("owner")
    if user["role"] != "admin" and owner != user["username"]:
        return jsonify({"ok": False, "error": "Solo propietario o administrador pueden gestionar permisos"}), 403
    perms = _load_permissions().get(project_id,{})
    users = _load_users()
    result=[]
    for username, permission in perms.items():
        result.append({"username":username,"permission":permission,"role":(users.get(username) or {}).get("role","")})
    return jsonify({"ok": True, "permissions": result})

@app.post("/projects/<project_id>/permissions")
def grant_project_permission(project_id):
    user, err = _require_auth()
    if err: return err
    path = _project_path(project_id)
    if not path.exists():
        return jsonify({"ok": False, "error": "Proyecto no encontrado"}), 404
    project = json.loads(path.read_text(encoding="utf-8"))
    owner = project.get("_institution",{}).get("owner")
    if user["role"] != "admin" and owner != user["username"]:
        return jsonify({"ok": False, "error": "Solo propietario o administrador pueden gestionar permisos"}), 403
    body=request.get_json(force=True)
    username=str(body.get("username","")).strip()
    permission=str(body.get("permission","viewer"))
    if permission not in ("viewer","editor"):
        return jsonify({"ok": False, "error": "Permiso inválido"}), 400
    if username not in _load_users():
        return jsonify({"ok": False, "error": "Usuario no configurado"}), 404
    perms=_load_permissions()
    perms.setdefault(project_id,{})[username]=permission
    _save_permissions(perms)
    _log_activity(project_id,user["username"],"Permiso concedido",f"{username}: {permission}")
    _notify(username,"Proyecto compartido",f"{user['username']} le concedió permiso {permission} sobre un proyecto.",project_id)
    return jsonify({"ok": True})

@app.delete("/projects/<project_id>/permissions/<username>")
def revoke_project_permission(project_id, username):
    user, err = _require_auth()
    if err: return err
    path = _project_path(project_id)
    if not path.exists():
        return jsonify({"ok": False, "error": "Proyecto no encontrado"}), 404
    project = json.loads(path.read_text(encoding="utf-8"))
    owner = project.get("_institution",{}).get("owner")
    if user["role"] != "admin" and owner != user["username"]:
        return jsonify({"ok": False, "error": "Solo propietario o administrador pueden gestionar permisos"}), 403
    perms=_load_permissions()
    if project_id in perms:
        perms[project_id].pop(username,None)
    _save_permissions(perms)
    _log_activity(project_id,user["username"],"Permiso revocado",username)
    _notify(username,"Permiso revocado",f"Se revocó su acceso a un proyecto.",project_id)
    return jsonify({"ok": True})

@app.get("/projects/<project_id>/versions")
def list_project_versions(project_id):
    user, err = _require_auth()
    if err: return err
    path = _project_path(project_id)
    if not path.exists():
        return jsonify({"ok": False, "error": "Proyecto no encontrado"}), 404
    project=json.loads(path.read_text(encoding="utf-8"))
    if not _can_read_project(user,project):
        return jsonify({"ok": False, "error": "Permisos insuficientes"}), 403
    out=[]
    for p in _version_project_dir(project_id).glob("*.json"):
        try:
            obj=json.loads(p.read_text(encoding="utf-8"))
            out.append({k:obj.get(k) for k in ("version","created_at","created_by","comment")})
        except Exception:
            pass
    out.sort(key=lambda x:int(x.get("version",0)),reverse=True)
    return jsonify({"ok": True, "versions": out})

@app.post("/projects/<project_id>/versions")
def create_project_version(project_id):
    user, err = _require_auth()
    if err: return err
    path=_project_path(project_id)
    if not path.exists():
        return jsonify({"ok": False, "error": "Proyecto no encontrado"}), 404
    current=json.loads(path.read_text(encoding="utf-8"))
    if not _can_edit_project(user,current):
        return jsonify({"ok": False, "error": "Permiso de edición requerido"}), 403
    body=request.get_json(force=True)
    project=body.get("project") or current
    comment=str(body.get("comment",""))
    v=_save_version(project_id,project,user["username"],comment)
    return jsonify({"ok": True, "version": v})

@app.get("/projects/<project_id>/versions/<version>")
def get_project_version(project_id,version):
    user, err = _require_auth()
    if err: return err
    path=_project_path(project_id)
    if not path.exists():
        return jsonify({"ok": False, "error": "Proyecto no encontrado"}),404
    current=json.loads(path.read_text(encoding="utf-8"))
    if not _can_read_project(user,current):
        return jsonify({"ok": False, "error": "Permisos insuficientes"}),403
    vpath=_version_project_dir(project_id)/f"{version}.json"
    if not vpath.exists():
        return jsonify({"ok": False, "error": "Versión no encontrada"}),404
    obj=json.loads(vpath.read_text(encoding="utf-8"))
    return jsonify({"ok": True, "project": obj.get("project"), "meta": {k:obj.get(k) for k in ("version","created_at","created_by","comment")}})

@app.post("/projects/<project_id>/restore/<version>")
def restore_project_version(project_id,version):
    user, err = _require_auth()
    if err: return err
    path=_project_path(project_id)
    if not path.exists():
        return jsonify({"ok": False, "error": "Proyecto no encontrado"}),404
    current=json.loads(path.read_text(encoding="utf-8"))
    if not _can_edit_project(user,current):
        return jsonify({"ok": False, "error": "Permiso de edición requerido"}),403
    vpath=_version_project_dir(project_id)/f"{version}.json"
    if not vpath.exists():
        return jsonify({"ok": False, "error": "Versión no encontrada"}),404
    old=json.loads(vpath.read_text(encoding="utf-8"))
    _save_version(project_id,current,user["username"],f"Backup automático antes de restaurar v{version}")
    restored=old.get("project") or {}
    restored["_institution"]=current.get("_institution",{})
    restored["_institution"]["updated_at"]=__import__("datetime").datetime.utcnow().isoformat(timespec="seconds")+"Z"
    restored["_institution"]["revision"]=int(current.get("_institution",{}).get("revision",1))+1
    path.write_text(json.dumps(restored,ensure_ascii=False,indent=2),encoding="utf-8")
    _log_activity(project_id,user["username"],"Versión restaurada",f"version {version}")
    return jsonify({"ok": True,"revision":restored["_institution"]["revision"]})

@app.get("/admin/summary")
def admin_summary():
    user, err = _require_auth(roles={"admin"})
    if err: return err
    users_cfg=_load_users()
    users=[{"username":u,"role":(v or {}).get("role","")} for u,v in users_cfg.items()]
    projects=list(PROJECT_DIR.glob("*.json"))
    versions_total=0
    for d in VERSION_DIR.glob("*"):
        if d.is_dir():
            versions_total += len(list(d.glob("*.json")))
    return jsonify({
        "ok": True,
        "users_total": len(users),
        "projects_total": len(projects),
        "versions_total": versions_total,
        "users": users
    })



# -----------------------------
# ValiStruct 2.5 comments / activity / notifications / conflict control
# -----------------------------
COMMENTS_FILE = PROJECT_DIR / "_comments.json"
ACTIVITY_FILE = PROJECT_DIR / "_activity.json"
NOTIFICATIONS_FILE = PROJECT_DIR / "_notifications.json"

def _read_json_file(path, default):
    try:
        return json.loads(path.read_text(encoding="utf-8")) if path.exists() else default
    except Exception:
        return default

def _atomic_write_text(path, text):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=path.name+".", suffix=".tmp", dir=str(path.parent))
    try:
        with os.fdopen(fd,"w",encoding="utf-8") as f:
            f.write(text)
            f.flush()
            os.fsync(f.fileno())
        os.replace(tmp,path)
    finally:
        if os.path.exists(tmp):
            try: os.remove(tmp)
            except Exception: pass

def _write_json_file(path, obj):
    try:
        _atomic_write_text(path,json.dumps(obj,ensure_ascii=False,indent=2))
    except Exception as e:
        raise RuntimeError(f"No se pudo escribir {Path(path).name}: {e}") from e

try:
    import portalocker
except Exception:
    portalocker = None

_LOCAL_LOCKS = {}
_LOCAL_LOCKS_GUARD = threading.Lock()

@contextmanager
def _file_lock(lock_path, timeout=15):
    lock_path = Path(lock_path)
    lock_path.parent.mkdir(parents=True,exist_ok=True)
    if portalocker is not None:
        with portalocker.Lock(str(lock_path),mode="a+",timeout=timeout):
            yield
    else:
        key=str(lock_path.resolve())
        with _LOCAL_LOCKS_GUARD:
            lock=_LOCAL_LOCKS.setdefault(key,threading.RLock())
        with lock:
            yield

def _mutate_json_file(path, default, mutator):
    lock_path = Path(str(path)+".lock")
    with _file_lock(lock_path):
        data = _read_json_file(path, default)
        result = mutator(data)
        _write_json_file(path, data)
        return result

def _log_activity(project_id, user, action, detail=""):
    def mut(data):
        data.setdefault(project_id,[]).append({
            "id":secrets.token_hex(8),"user":user,"action":action,"detail":detail,
            "created_at":__import__("datetime").datetime.utcnow().isoformat(timespec="seconds")+"Z"
        })
        data[project_id]=data[project_id][-1000:]
    _mutate_json_file(ACTIVITY_FILE,{},mut)

def _notify(username,title,message,project_id=None):
    def mut(data):
        data.setdefault(username,[]).append({
            "id":secrets.token_hex(8),"title":title,"message":message,
            "project_id":project_id,"read":False,
            "created_at":__import__("datetime").datetime.utcnow().isoformat(timespec="seconds")+"Z"
        })
        data[username]=data[username][-500:]
    _mutate_json_file(NOTIFICATIONS_FILE,{},mut)

def _project_collaborators(project_id):
    perms=_load_permissions().get(project_id,{})
    return list(perms.keys())

@app.get("/projects/<project_id>/comments")
def list_project_comments(project_id):
    user, err=_require_auth()
    if err: return err
    path=_project_path(project_id)
    if not path.exists():
        return jsonify({"ok":False,"error":"Proyecto no encontrado"}),404
    project=json.loads(path.read_text(encoding="utf-8"))
    if not _can_read_project(user,project):
        return jsonify({"ok":False,"error":"Permisos insuficientes"}),403
    data=_read_json_file(COMMENTS_FILE,{})
    comments=data.get(project_id,[])
    return jsonify({"ok":True,"comments":comments})

@app.post("/projects/<project_id>/comments")
def add_project_comment(project_id):
    user, err=_require_auth()
    if err: return err
    path=_project_path(project_id)
    if not path.exists():
        return jsonify({"ok":False,"error":"Proyecto no encontrado"}),404
    project=json.loads(path.read_text(encoding="utf-8"))
    if not _can_read_project(user,project):
        return jsonify({"ok":False,"error":"Permisos insuficientes"}),403
    body=request.get_json(force=True)
    text=str(body.get("text","")).strip()
    ctype=str(body.get("type","general")).strip()
    if not text:
        return jsonify({"ok":False,"error":"Comentario vacío"}),400
    data=_read_json_file(COMMENTS_FILE,{})
    comment={
        "id":secrets.token_hex(8),
        "author":user["username"],
        "type":ctype,
        "text":text,
        "resolved":False,
        "created_at":__import__("datetime").datetime.utcnow().isoformat(timespec="seconds")+"Z"
    }
    data.setdefault(project_id,[]).append(comment)
    _write_json_file(COMMENTS_FILE,data)
    _log_activity(project_id,user["username"],"Comentario agregado",ctype)
    owner=project.get("_institution",{}).get("owner")
    recipients=set(_project_collaborators(project_id)+([owner] if owner else []))
    recipients.discard(user["username"])
    import re as _re
    mentions=set(_re.findall(r"@([A-Za-z0-9_.-]+)",text))
    for u in recipients:
        _notify(u,"Nuevo comentario",f"{user['username']} agregó un comentario {ctype}.",project_id)
    for u in mentions:
        if u in _load_users() and u != user["username"]:
            _notify(u,"Mención en comentario",f"{user['username']} le mencionó en un comentario.",project_id)
    return jsonify({"ok":True,"comment":comment,"mentions":sorted(mentions)})

@app.post("/projects/<project_id>/comments/<comment_id>/toggle")
def toggle_project_comment(project_id,comment_id):
    user, err=_require_auth()
    if err: return err
    path=_project_path(project_id)
    if not path.exists():
        return jsonify({"ok":False,"error":"Proyecto no encontrado"}),404
    project=json.loads(path.read_text(encoding="utf-8"))
    if not _can_edit_project(user,project):
        return jsonify({"ok":False,"error":"Permiso de edición requerido"}),403
    data=_read_json_file(COMMENTS_FILE,{})
    found=None
    for c in data.get(project_id,[]):
        if c.get("id")==comment_id:
            c["resolved"]=not bool(c.get("resolved"))
            c["resolved_by"]=user["username"]
            found=c
            break
    if not found:
        return jsonify({"ok":False,"error":"Comentario no encontrado"}),404
    _write_json_file(COMMENTS_FILE,data)
    _log_activity(project_id,user["username"],"Comentario actualizado","resuelto" if found["resolved"] else "reabierto")
    return jsonify({"ok":True,"comment":found})

@app.get("/projects/<project_id>/activity")
def project_activity(project_id):
    user, err=_require_auth()
    if err: return err
    path=_project_path(project_id)
    if not path.exists():
        return jsonify({"ok":False,"error":"Proyecto no encontrado"}),404
    project=json.loads(path.read_text(encoding="utf-8"))
    if not _can_read_project(user,project):
        return jsonify({"ok":False,"error":"Permisos insuficientes"}),403
    data=_read_json_file(ACTIVITY_FILE,{})
    return jsonify({"ok":True,"activity":list(reversed(data.get(project_id,[])))})

@app.get("/notifications")
def get_notifications():
    user, err=_require_auth()
    if err: return err
    data=_read_json_file(NOTIFICATIONS_FILE,{})
    items=list(reversed(data.get(user["username"],[])))
    return jsonify({"ok":True,"notifications":items})

@app.post("/notifications/read-all")
def read_all_notifications():
    user, err=_require_auth()
    if err: return err
    data=_read_json_file(NOTIFICATIONS_FILE,{})
    for n in data.get(user["username"],[]):
        n["read"]=True
    _write_json_file(NOTIFICATIONS_FILE,data)
    return jsonify({"ok":True})



# -----------------------------
# ValiStruct 2.6 tasks / review status / mentions / telemetry
# -----------------------------
TASKS_FILE = PROJECT_DIR / "_tasks.json"
REVIEW_STATUS_FILE = PROJECT_DIR / "_review_status.json"
TELEMETRY_FILE = PROJECT_DIR / "_telemetry.json"

@app.get("/projects/<project_id>/tasks")
def list_project_tasks(project_id):
    user, err=_require_auth()
    if err: return err
    path=_project_path(project_id)
    if not path.exists(): return jsonify({"ok":False,"error":"Proyecto no encontrado"}),404
    project=json.loads(path.read_text(encoding="utf-8"))
    if not _can_read_project(user,project): return jsonify({"ok":False,"error":"Permisos insuficientes"}),403
    tasks=_read_json_file(TASKS_FILE,{}).get(project_id,[])
    status=_read_json_file(REVIEW_STATUS_FILE,{}).get(project_id,"pending")
    return jsonify({"ok":True,"tasks":tasks,"review_status":status})

@app.post("/projects/<project_id>/tasks")
def add_project_task(project_id):
    user, err=_require_auth()
    if err: return err
    path=_project_path(project_id)
    if not path.exists(): return jsonify({"ok":False,"error":"Proyecto no encontrado"}),404
    project=json.loads(path.read_text(encoding="utf-8"))
    if not _can_edit_project(user,project): return jsonify({"ok":False,"error":"Permiso de edición requerido"}),403
    body=request.get_json(force=True)
    text=str(body.get("text","")).strip()
    assignee=str(body.get("assignee","")).strip()
    priority=str(body.get("priority","medium"))
    if not text or not assignee: return jsonify({"ok":False,"error":"Tarea y usuario son obligatorios"}),400
    if assignee not in _load_users(): return jsonify({"ok":False,"error":"Usuario no configurado"}),404
    data=_read_json_file(TASKS_FILE,{})
    task={"id":secrets.token_hex(8),"text":text,"assignee":assignee,"priority":priority,"done":False,
          "created_by":user["username"],"created_at":__import__("datetime").datetime.utcnow().isoformat(timespec="seconds")+"Z"}
    data.setdefault(project_id,[]).append(task);_write_json_file(TASKS_FILE,data)
    _log_activity(project_id,user["username"],"Tarea asignada",f"{assignee}: {text[:120]}")
    _notify(assignee,"Nueva tarea",text[:240],project_id)
    return jsonify({"ok":True,"task":task})

@app.post("/projects/<project_id>/tasks/<task_id>/toggle")
def toggle_project_task(project_id,task_id):
    user, err=_require_auth()
    if err: return err
    path=_project_path(project_id)
    if not path.exists():
        return jsonify({"ok":False,"error":"Proyecto no encontrado"}),404
    project=json.loads(path.read_text(encoding="utf-8"))
    data=_read_json_file(TASKS_FILE,{})
    found=None
    for t in data.get(project_id,[]):
        if t.get("id")==task_id:
            allowed = (
                user["role"]=="admin" or
                user["username"] in (t.get("assignee"),t.get("created_by")) or
                _can_edit_project(user,project)
            )
            if not allowed:
                return jsonify({"ok":False,"error":"Permisos insuficientes"}),403
            t["done"]=not bool(t.get("done"))
            t["status"]="done" if t["done"] else "todo"
            found=t
            break
    if not found:
        return jsonify({"ok":False,"error":"Tarea no encontrada"}),404
    try:
        _write_json_file(TASKS_FILE,data)
    except RuntimeError as e:
        return jsonify({"ok":False,"error":str(e)}),500
    _log_activity(project_id,user["username"],"Tarea actualizada","completada" if found["done"] else "reabierta")
    return jsonify({"ok":True,"task":found})

@app.put("/projects/<project_id>/review-status")
def set_review_status(project_id):
    user, err=_require_auth()
    if err:return err
    path=_project_path(project_id)
    if not path.exists():return jsonify({"ok":False,"error":"Proyecto no encontrado"}),404
    project=json.loads(path.read_text(encoding="utf-8"))
    if not _can_edit_project(user,project):return jsonify({"ok":False,"error":"Permiso de edición requerido"}),403
    status=str(request.get_json(force=True).get("status","pending"))
    if status not in ("pending","in_review","approved"):return jsonify({"ok":False,"error":"Estado inválido"}),400
    data=_read_json_file(REVIEW_STATUS_FILE,{})
    data[project_id]=status;_write_json_file(REVIEW_STATUS_FILE,data)
    _log_activity(project_id,user["username"],"Estado de revisión",status)
    return jsonify({"ok":True,"status":status})

@app.post("/collaboration-docx")
def collaboration_docx():
    _user, _auth_err = _compute_auth_guard()
    if _auth_err: return _auth_err
    from io import BytesIO
    from flask import send_file
    try:
        from docx import Document
        from docx.shared import Pt
    except Exception as e:
        return f"python-docx no disponible: {e}",500
    body=request.get_json(force=True)
    doc=Document()
    doc.styles["Normal"].font.name="Arial";doc.styles["Normal"].font.size=Pt(11)
    doc.add_heading("ValiStruct · Bitácora colaborativa",0)
    doc.add_paragraph(f"Estado de revisión: {body.get('review_status','pending')}")
    doc.add_heading("Actividad",level=1)
    for x in body.get("activity",[]):
        doc.add_paragraph(f"{x.get('created_at','')} · {x.get('user','')} · {x.get('action','')} · {x.get('detail','')}")
    doc.add_heading("Comentarios",level=1)
    for x in body.get("comments",[]):
        doc.add_paragraph(f"{x.get('created_at','')} · {x.get('author','')} · {x.get('type','')} · {x.get('text','')} · {'Resuelto' if x.get('resolved') else 'Abierto'}")
    doc.add_heading("Tareas",level=1)
    for x in body.get("tasks",[]):
        doc.add_paragraph(f"{x.get('assignee','')} · {x.get('priority','')} · {x.get('text','')} · {'Completada' if x.get('done') else 'Pendiente'}")
    bio=BytesIO();doc.save(bio);bio.seek(0)
    return send_file(bio,mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                     as_attachment=True,download_name="ValiStruct_bitacora_colaborativa.docx")

@app.post("/telemetry")
def telemetry():
    body=request.get_json(force=True)
    event=str(body.get("event","")).strip()
    version=str(body.get("app_version","")).strip()
    allowed={"button_click","frontend_error","backend_error","analysis_run","export_run"}
    if event not in allowed:return jsonify({"ok":False,"error":"Evento no permitido"}),400
    data=_read_json_file(TELEMETRY_FILE,{"events":{}})
    key=f"{version}:{event}"
    data["events"][key]=int(data["events"].get(key,0))+1
    data["updated_at"]=__import__("datetime").datetime.utcnow().isoformat(timespec="seconds")+"Z"
    _write_json_file(TELEMETRY_FILE,data)
    return jsonify({"ok":True})



# -----------------------------
# ValiStruct 2.7 kanban / approvals / releases / audit / monitor
# -----------------------------
APPROVALS_FILE = PROJECT_DIR / "_approvals.json"
RELEASES_FILE = PROJECT_DIR / "_releases.json"

@app.put("/projects/<project_id>/tasks/<task_id>/status")
def set_task_status(project_id,task_id):
    user, err=_require_auth()
    if err:return err
    path=_project_path(project_id)
    if not path.exists():return jsonify({"ok":False,"error":"Proyecto no encontrado"}),404
    project=json.loads(path.read_text(encoding="utf-8"))
    if not _can_edit_project(user,project):return jsonify({"ok":False,"error":"Permiso de edición requerido"}),403
    status=str(request.get_json(force=True).get("status","todo"))
    if status not in ("todo","doing","done"):return jsonify({"ok":False,"error":"Estado inválido"}),400
    data=_read_json_file(TASKS_FILE,{})
    found=None
    for t in data.get(project_id,[]):
        if t.get("id")==task_id:
            t["status"]=status
            t["done"]=status=="done"
            found=t;break
    if not found:return jsonify({"ok":False,"error":"Tarea no encontrada"}),404
    _write_json_file(TASKS_FILE,data)
    _log_activity(project_id,user["username"],"Kanban actualizado",f"{task_id}: {status}")
    return jsonify({"ok":True,"task":found})

@app.get("/projects/<project_id>/approvals")
def list_approvals(project_id):
    user, err=_require_auth()
    if err:return err
    path=_project_path(project_id)
    if not path.exists():return jsonify({"ok":False,"error":"Proyecto no encontrado"}),404
    project=json.loads(path.read_text(encoding="utf-8"))
    if not _can_read_project(user,project):return jsonify({"ok":False,"error":"Permisos insuficientes"}),403
    data=_read_json_file(APPROVALS_FILE,{})
    return jsonify({"ok":True,"approvals":list(reversed(data.get(project_id,[])))})

@app.post("/projects/<project_id>/approvals")
def add_approval(project_id):
    user, err=_require_auth()
    if err:return err
    path=_project_path(project_id)
    if not path.exists():return jsonify({"ok":False,"error":"Proyecto no encontrado"}),404
    project=json.loads(path.read_text(encoding="utf-8"))
    if not _can_edit_project(user,project):return jsonify({"ok":False,"error":"Permiso de edición requerido"}),403
    body=request.get_json(force=True)
    typ=str(body.get("type","methodology"))
    decision=str(body.get("decision","approved"))
    comment=str(body.get("comment",""))
    created=__import__("datetime").datetime.utcnow().isoformat(timespec="seconds")+"Z"
    base=f"{project_id}|{user['username']}|{typ}|{decision}|{comment}|{created}"
    signature_hash=hashlib.sha256(base.encode("utf-8")).hexdigest()
    item={"id":secrets.token_hex(8),"type":typ,"decision":decision,"comment":comment,
          "user":user["username"],"created_at":created,"signature_hash":signature_hash}
    data=_read_json_file(APPROVALS_FILE,{})
    data.setdefault(project_id,[]).append(item);_write_json_file(APPROVALS_FILE,data)
    _log_activity(project_id,user["username"],"Aprobación registrada",f"{typ}: {decision}")
    return jsonify({"ok":True,"approval":item})

@app.get("/projects/<project_id>/releases")
def list_releases(project_id):
    user, err=_require_auth()
    if err:return err
    path=_project_path(project_id)
    if not path.exists():return jsonify({"ok":False,"error":"Proyecto no encontrado"}),404
    project=json.loads(path.read_text(encoding="utf-8"))
    if not _can_read_project(user,project):return jsonify({"ok":False,"error":"Permisos insuficientes"}),403
    data=_read_json_file(RELEASES_FILE,{})
    return jsonify({"ok":True,"releases":list(reversed(data.get(project_id,[])))})

@app.post("/projects/<project_id>/releases")
def create_release(project_id):
    user, err=_require_auth()
    if err:return err
    path=_project_path(project_id)
    if not path.exists():return jsonify({"ok":False,"error":"Proyecto no encontrado"}),404
    project=json.loads(path.read_text(encoding="utf-8"))
    if not _can_edit_project(user,project):return jsonify({"ok":False,"error":"Permiso de edición requerido"}),403
    body=request.get_json(force=True)
    tag=str(body.get("tag","")).strip()
    if not tag:return jsonify({"ok":False,"error":"Etiqueta requerida"}),400
    data=_read_json_file(RELEASES_FILE,{})
    if any(x.get("tag")==tag for x in data.get(project_id,[])):
        return jsonify({"ok":False,"error":"La etiqueta ya existe"}),409
    version=_save_version(project_id,project,user["username"],f"Release {tag}")
    item={"id":secrets.token_hex(8),"tag":tag,"type":str(body.get("type","milestone")),
          "notes":str(body.get("notes","")),"version":version,"created_by":user["username"],
          "created_at":__import__("datetime").datetime.utcnow().isoformat(timespec="seconds")+"Z"}
    data.setdefault(project_id,[]).append(item);_write_json_file(RELEASES_FILE,data)
    _log_activity(project_id,user["username"],"Release creado",tag)
    return jsonify({"ok":True,"release":item})

def _project_audit_data(project_id,user):
    path=_project_path(project_id)
    if not path.exists():raise FileNotFoundError("Proyecto no encontrado")
    project=json.loads(path.read_text(encoding="utf-8"))
    if not _can_read_project(user,project):raise PermissionError("Permisos insuficientes")
    activity=_read_json_file(ACTIVITY_FILE,{}).get(project_id,[])
    comments=_read_json_file(COMMENTS_FILE,{}).get(project_id,[])
    tasks=_read_json_file(TASKS_FILE,{}).get(project_id,[])
    approvals=_read_json_file(APPROVALS_FILE,{}).get(project_id,[])
    releases=_read_json_file(RELEASES_FILE,{}).get(project_id,[])
    review_status=_read_json_file(REVIEW_STATUS_FILE,{}).get(project_id,"pending")
    versions=[]
    for p in _version_project_dir(project_id).glob("*.json"):
        try:
            obj=json.loads(p.read_text(encoding="utf-8"))
            versions.append({k:obj.get(k) for k in ("version","created_at","created_by","comment")})
        except Exception:pass
    return {
        "project":project,"activity":activity,"comments":comments,"tasks":tasks,
        "approvals":approvals,"releases":releases,"review_status":review_status,"versions":versions
    }

@app.get("/projects/<project_id>/audit-summary")
def audit_summary(project_id):
    user,err=_require_auth()
    if err:return err
    try:data=_project_audit_data(project_id,user)
    except FileNotFoundError:return jsonify({"ok":False,"error":"Proyecto no encontrado"}),404
    except PermissionError:return jsonify({"ok":False,"error":"Permisos insuficientes"}),403
    return jsonify({"ok":True,"project_name":data["project"].get("name","Proyecto"),
                    "review_status":data["review_status"],
                    "counts":{"activity":len(data["activity"]),"comments":len(data["comments"]),
                              "tasks":len(data["tasks"]),"approvals":len(data["approvals"]),
                              "releases":len(data["releases"]),"versions":len(data["versions"])}})

@app.get("/projects/<project_id>/audit-package")
def audit_package(project_id):
    from io import BytesIO
    import zipfile as _zipfile
    user,err=_require_auth()
    if err:return err
    try:data=_project_audit_data(project_id,user)
    except FileNotFoundError:return jsonify({"ok":False,"error":"Proyecto no encontrado"}),404
    except PermissionError:return jsonify({"ok":False,"error":"Permisos insuficientes"}),403
    bio=BytesIO()
    with _zipfile.ZipFile(bio,"w",_zipfile.ZIP_DEFLATED) as z:
        z.writestr("project.json",json.dumps(data["project"],ensure_ascii=False,indent=2))
        z.writestr("activity.json",json.dumps(data["activity"],ensure_ascii=False,indent=2))
        z.writestr("comments.json",json.dumps(data["comments"],ensure_ascii=False,indent=2))
        z.writestr("tasks.json",json.dumps(data["tasks"],ensure_ascii=False,indent=2))
        z.writestr("approvals.json",json.dumps(data["approvals"],ensure_ascii=False,indent=2))
        z.writestr("releases.json",json.dumps(data["releases"],ensure_ascii=False,indent=2))
        z.writestr("versions.json",json.dumps(data["versions"],ensure_ascii=False,indent=2))
        z.writestr("review_status.txt",str(data["review_status"]))
        manifest={"app":"ValiStruct","version":"2.7","generated_at":__import__("datetime").datetime.utcnow().isoformat(timespec="seconds")+"Z"}
        z.writestr("manifest.json",json.dumps(manifest,ensure_ascii=False,indent=2))
    bio.seek(0)
    return send_file(bio,mimetype="application/zip",as_attachment=True,download_name="ValiStruct_paquete_auditoria.zip")

@app.get("/monitor")
def monitor():
    user,err=_require_auth()
    if err:return err
    if AUTH_ENABLED and user.get("role") not in ("admin","researcher","teacher"):
        return jsonify({"ok":False,"error":"Permisos insuficientes"}),403
    uptime=max(0,int(time.time()-APP_STARTED_AT))
    h,m=divmod(uptime,3600);m,s=divmod(m,60)
    projects_total=len([p for p in PROJECT_DIR.glob("*.json") if not p.name.startswith("_")])
    comments=sum(len(v) for v in _read_json_file(COMMENTS_FILE,{}).values())
    tasks=sum(len(v) for v in _read_json_file(TASKS_FILE,{}).values())
    notifications=sum(len(v) for v in _read_json_file(NOTIFICATIONS_FILE,{}).values())
    return jsonify({"ok":True,"status":"operativo","app_version":"3.0.0-rc.6",
                    "uptime_seconds":uptime,"uptime_human":f"{h}h {m}m {s}s",
                    "projects_total":projects_total,"comments_total":comments,
                    "tasks_total":tasks,"notifications_total":notifications,
                    "incidents_total":len(_read_json_file(INCIDENTS_FILE,[])),
                    "r_available":bool(find_rscript()),"auth_enabled":AUTH_ENABLED,
                    "project_library_enabled":PROJECT_LIBRARY_ENABLED})



# -----------------------------
# ValiStruct 2.8 backup / restore / incidents
# -----------------------------
INCIDENTS_FILE = PROJECT_DIR / "_incidents.json"

@app.get("/admin/backup")
def admin_backup():
    from io import BytesIO
    import zipfile as _zipfile
    user,err=_require_auth(roles={"admin"})
    if err:return err
    bio=BytesIO()
    with _zipfile.ZipFile(bio,"w",_zipfile.ZIP_DEFLATED) as z:
        for p in PROJECT_DIR.rglob("*"):
            if p.is_file():
                z.write(p,arcname=str(p.relative_to(PROJECT_DIR)))
        manifest={
            "app":"ValiStruct","version":"2.8",
            "generated_at":__import__("datetime").datetime.utcnow().isoformat(timespec="seconds")+"Z",
            "generated_by":user["username"]
        }
        z.writestr("_backup_manifest.json",json.dumps(manifest,ensure_ascii=False,indent=2))
    bio.seek(0)
    _log_activity("_system",user["username"],"Respaldo institucional creado","")
    return send_file(bio,mimetype="application/zip",as_attachment=True,download_name="ValiStruct_respaldo_institucional.zip")

@app.post("/admin/restore")
def admin_restore():
    import zipfile as _zipfile
    import tempfile as _tempfile
    user,err=_require_auth(roles={"admin"})
    if err:return err
    if "file" not in request.files:
        return jsonify({"ok":False,"error":"No se recibió respaldo"}),400
    f=request.files["file"]
    tmp=None
    try:
        fd,tmp=_tempfile.mkstemp(suffix=".zip");os.close(fd);f.save(tmp)
        restored=0
        with _zipfile.ZipFile(tmp,"r") as z:
            for info in z.infolist():
                name=info.filename
                if name.startswith("/") or ".." in Path(name).parts or name=="_backup_manifest.json":
                    continue
                target=(PROJECT_DIR / name).resolve()
                if PROJECT_DIR.resolve() not in target.parents and target != PROJECT_DIR.resolve():
                    continue
                if info.is_dir():
                    target.mkdir(parents=True,exist_ok=True);continue
                target.parent.mkdir(parents=True,exist_ok=True)
                with z.open(info,"r") as src, open(target,"wb") as dst:
                    dst.write(src.read())
                restored+=1
        _log_activity("_system",user["username"],"Respaldo institucional restaurado",f"{restored} archivos")
        return jsonify({"ok":True,"files_restored":restored})
    except Exception as e:
        return jsonify({"ok":False,"error":str(e)}),400
    finally:
        if tmp and os.path.exists(tmp):
            try:os.remove(tmp)
            except Exception:pass

@app.get("/incidents")
def list_incidents():
    user,err=_require_auth()
    if err:return err
    data=_read_json_file(INCIDENTS_FILE,[])
    if user.get("role")=="admin":
        items=data
    else:
        items=[x for x in data if x.get("created_by")==user.get("username")]
    items=sorted(items,key=lambda x:x.get("created_at",""),reverse=True)
    return jsonify({"ok":True,"incidents":items})

@app.post("/incidents")
def create_incident():
    user,err=_require_auth()
    if err:return err
    body=request.get_json(force=True)
    description=str(body.get("description","")).strip()
    if not description:return jsonify({"ok":False,"error":"Descripción requerida"}),400
    item={
        "id":secrets.token_hex(8),
        "type":str(body.get("type","bug")),
        "severity":str(body.get("severity","medium")),
        "description":description[:4000],
        "status":"open",
        "created_by":user["username"],
        "created_at":__import__("datetime").datetime.utcnow().isoformat(timespec="seconds")+"Z"
    }
    data=_read_json_file(INCIDENTS_FILE,[])
    data.append(item);_write_json_file(INCIDENTS_FILE,data[-2000:])
    return jsonify({"ok":True,"incident":item})

@app.post("/incidents/<incident_id>/toggle")
def toggle_incident(incident_id):
    user,err=_require_auth()
    if err:return err
    data=_read_json_file(INCIDENTS_FILE,[])
    found=None
    for item in data:
        if item.get("id")==incident_id:
            if user.get("role")!="admin" and item.get("created_by")!=user.get("username"):
                return jsonify({"ok":False,"error":"Permisos insuficientes"}),403
            item["status"]="open" if item.get("status")=="closed" else "closed"
            item["updated_at"]=__import__("datetime").datetime.utcnow().isoformat(timespec="seconds")+"Z"
            found=item;break
    if not found:return jsonify({"ok":False,"error":"Incidencia no encontrada"}),404
    _write_json_file(INCIDENTS_FILE,data)
    return jsonify({"ok":True,"incident":found})



# -----------------------------
# ValiStruct 3.0 RC1 release-candidate checks
# -----------------------------
@app.get("/version")
def version_info():
    return jsonify({
        "ok": True,
        "app": "ValiStruct",
        "version": "3.0.0-rc.6",
        "project_format": "3.0",
        "release_channel": "release-candidate"
    })

@app.get("/rc-check")
def rc_check():
    checks=[]
    def add(name,status,detail=""):
        checks.append({"name":name,"status":status,"detail":detail})

    add("Python API","pass","Flask activo.")
    add("Project directory","pass" if PROJECT_DIR.exists() else "fail",str(PROJECT_DIR))
    add("Project directory writable","pass" if os.access(PROJECT_DIR,os.W_OK) else "fail","Escritura requerida.")
    add("Rscript","pass" if find_rscript() else "warn","Rscript disponible." if find_rscript() else "Rscript no localizado.")
    if find_rscript():
        try:
            rp=subprocess.run([find_rscript(),"-e",'cat(requireNamespace("lavaan",quietly=TRUE))'],
                              capture_output=True,text=True,timeout=20)
            add("R package lavaan","pass" if "TRUE" in rp.stdout else "fail",
                "Disponible." if "TRUE" in rp.stdout else "No disponible.")
        except Exception as e:
            add("R package lavaan","warn",str(e))
    try:
        import openpyxl
        add("openpyxl","pass",getattr(openpyxl,"__version__",""))
    except Exception as e:
        add("openpyxl","fail",str(e))
    try:
        import pyreadstat
        add("pyreadstat","pass",getattr(pyreadstat,"__version__",""))
    except Exception as e:
        add("pyreadstat","fail",str(e))
    try:
        from docx import Document
        add("python-docx","pass","Disponible.")
    except Exception as e:
        add("python-docx","fail",str(e))
    add("Auth configuration","pass" if AUTH_ENABLED else "warn","Activa." if AUTH_ENABLED else "Desactivada.")
    add("Institution library","pass" if PROJECT_LIBRARY_ENABLED else "warn","Activa." if PROJECT_LIBRARY_ENABLED else "Desactivada.")
    return jsonify({"ok":True,"release":"3.0.0-rc.6","checks":checks})

@app.get("/security-status")
def security_status():
    checks=[]
    def add(name,status,detail=""):
        checks.append({"name":name,"status":status,"detail":detail})
    origins=os.environ.get("VALISTRUCT_ALLOWED_ORIGINS","").strip()
    env_mode=os.environ.get("VALISTRUCT_ENV","development").strip().lower()
    add("Upload limit","pass" if app.config.get("MAX_CONTENT_LENGTH") else "warn",
        f"{int(app.config.get('MAX_CONTENT_LENGTH',0)/1024/1024)} MB" if app.config.get("MAX_CONTENT_LENGTH") else "Sin límite")
    add("CORS restricted","pass" if origins and "*" not in origins else ("warn" if env_mode=="development" else "fail"),
        origins or f"No definido en modo {env_mode}")
    add("Auth","pass" if AUTH_ENABLED else "warn","Activada" if AUTH_ENABLED else "Desactivada")
    add("Project library","pass" if PROJECT_LIBRARY_ENABLED else "warn","Activada" if PROJECT_LIBRARY_ENABLED else "Desactivada")
    add("Security headers","pass","X-Content-Type-Options, X-Frame-Options, Referrer-Policy, no-store.")
    add("Backup script","pass" if (Path(HERE).parent/"deployment"/"backup.sh").exists() else "warn","Script de respaldo.")
    return jsonify({"ok":True,"checks":checks})



@app.get("/runtime-audit")
def runtime_audit():
    """Lightweight runtime checks for RC deployment readiness."""
    checks = []
    def add(name, status, detail=""):
        checks.append({"name": name, "status": status, "detail": detail})

    # Core imports / globals
    add("send_file", "pass" if callable(send_file) else "fail", "Flask send_file import")
    add("PROJECT_DIR", "pass" if PROJECT_DIR.exists() else "fail", str(PROJECT_DIR))
    add("PROJECT_DIR writable", "pass" if os.access(PROJECT_DIR, os.W_OK) else "fail", "Write permission")
    add("Rscript", "pass" if find_rscript() else "warn", find_rscript() or "not found")

    # File-backed stores
    stores = [
        ("permissions", PERMISSION_FILE),
        ("comments", COMMENTS_FILE),
        ("activity", ACTIVITY_FILE),
        ("notifications", NOTIFICATIONS_FILE),
        ("tasks", TASKS_FILE),
        ("review_status", REVIEW_STATUS_FILE),
        ("approvals", APPROVALS_FILE),
        ("releases", RELEASES_FILE),
        ("incidents", INCIDENTS_FILE),
    ]
    for label, path in stores:
        parent = path.parent
        add(f"store:{label}", "pass" if parent.exists() and os.access(parent, os.W_OK) else "fail", str(path))

    # Required Python dependencies
    modules = [
        ("openpyxl", "openpyxl"),
        ("pyreadstat", "pyreadstat"),
        ("python-docx", "docx"),
        ("xlsxwriter", "xlsxwriter"),
    ]
    import importlib.util
    for label, module in modules:
        add(f"dependency:{label}", "pass" if importlib.util.find_spec(module) else "warn", module)

    return jsonify({"ok": True, "release": "3.0.0-rc.6", "checks": checks})

if __name__ == "__main__":
    app.run(host="127.0.0.1", port=8765, debug=False)
