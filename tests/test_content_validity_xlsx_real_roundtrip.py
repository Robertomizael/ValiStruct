from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCHEMA = "content-validity-xlsx-v1"


def _make_book(profile: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "ValiStruct"

    if profile == "cvi":
        judges, items = 4, 2
        rows = [
            ["VALISTRUCT · VALIDACIÓN DE CONTENIDO POR EXPERTOS"],
            ["VALISTRUCT_PROFILE", "cvi"],
            ["VALISTRUCT_SCHEMA", SCHEMA],
            ["Número de jueces", judges],
            ["Número de ítems", items],
            ["Escala mínima", 1],
            ["Escala máxima", 4],
            ["Relevante desde", 3],
            [],
            ["Dimensión", "No. ítem", "Ítem", "Juez 1", "Juez 2", "Juez 3", "Juez 4", "Observaciones"],
            ["D1", 1, "Ítem 1", 4, 4, 3, 4, ""],
            ["D1", 2, "Ítem 2", 3, 3, 2, 4, "Revisar redacción"],
            [],
            ["IDENTIFICACIÓN DE EXPERTOS"],
            ["Columna", "Nombre o código del experto", "Correo/identificador opcional"],
            ["Juez 1", "E01", ""],
            ["Juez 2", "E02", ""],
            ["Juez 3", "E03", ""],
            ["Juez 4", "E04", ""],
        ]
    elif profile == "lawshe":
        judges, items = 4, 2
        rows = [
            ["VALISTRUCT · VALIDACIÓN DE CONTENIDO POR EXPERTOS"],
            ["VALISTRUCT_PROFILE", "lawshe"],
            ["VALISTRUCT_SCHEMA", SCHEMA],
            ["Número de jueces", judges],
            ["Número de ítems", items],
            ["Opciones válidas", "Esencial | Útil pero no esencial | No necesario"],
            [],
            ["Dimensión", "No. ítem", "Ítem", "Juez 1", "Juez 2", "Juez 3", "Juez 4", "Observaciones"],
            ["D1", 1, "Ítem 1", "Esencial", "Esencial", "Esencial", "Esencial", ""],
            ["D1", 2, "Ítem 2", "Esencial", "Útil pero no esencial", "Esencial", "No necesario", ""],
            [],
            ["IDENTIFICACIÓN DE EXPERTOS"],
            ["Columna", "Nombre o código del experto", "Correo/identificador opcional"],
            ["Juez 1", "E01", ""],
            ["Juez 2", "E02", ""],
            ["Juez 3", "E03", ""],
            ["Juez 4", "E04", ""],
        ]
    elif profile == "delphi":
        judges, items = 4, 2
        rows = [
            ["VALISTRUCT · VALIDACIÓN DE CONTENIDO POR EXPERTOS"],
            ["VALISTRUCT_PROFILE", "delphi"],
            ["VALISTRUCT_SCHEMA", SCHEMA],
            ["Modalidad", "Delphi modificado"],
            ["Ronda", 1],
            ["Número de expertos", judges],
            ["Número de ítems", items],
            ["Escala mínima", 1],
            ["Escala máxima", 5],
            [],
            ["Dimensión", "No. ítem", "Ítem", "Ronda", "Experto 1", "Experto 2", "Experto 3", "Experto 4", "Observaciones"],
            ["D1", 1, "Ítem 1", 1, 5, 4, 4, 5, ""],
            ["D1", 2, "Ítem 2", 1, 4, 4, 3, 4, "Ajustar redacción"],
            [],
            ["IDENTIFICACIÓN DE EXPERTOS"],
            ["Columna", "Nombre o código del experto", "Correo/identificador opcional"],
            ["Experto 1", "E01", ""],
            ["Experto 2", "E02", ""],
            ["Experto 3", "E03", ""],
            ["Experto 4", "E04", ""],
        ]
    else:
        raise AssertionError(profile)

    for row in rows:
        ws.append(row)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue(), judges, items


def _metadata_and_matrix(data: bytes):
    wb = load_workbook(BytesIO(data), data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    metadata = {}
    for row in rows[:30]:
        if len(row) >= 2 and row[0] not in (None, ""):
            metadata[str(row[0]).strip()] = row[1]

    header_index = next(
        i for i, row in enumerate(rows)
        if "Ítem" in row and any(str(v or "").startswith(("Juez ", "Experto ")) for v in row)
    )
    headers = [str(v or "").strip() for v in rows[header_index]]
    expert_headers = [h for h in headers if h.startswith("Juez ") or h.startswith("Experto ")]
    item_index = headers.index("Ítem")
    body = []
    for row in rows[header_index + 1:]:
        item = row[item_index] if item_index < len(row) else None
        if item in (None, ""):
            break
        body.append(row)
    return metadata, expert_headers, body, rows


def test_active_renderer_declares_current_schema_and_three_profiles():
    src = (ROOT / "desktop" / "content-validity-xlsx-hardening.js").read_text(encoding="utf-8")
    assert f"const SCHEMA='{SCHEMA}'" in src
    for token in ("cviDownloadXlsx:'cvi'", "lawsheDownloadXlsx:'lawshe'", "delphiDownloadXlsx:'delphi'"):
        assert token in src
    assert "VALISTRUCT_PROFILE" in src
    assert "VALISTRUCT_SCHEMA" in src
    assert "IDENTIFICACIÓN DE EXPERTOS" in src


def test_real_xlsx_roundtrip_preserves_profile_schema_matrix_and_experts():
    for profile in ("cvi", "lawshe", "delphi"):
        data, judges, items = _make_book(profile)
        assert data[:2] == b"PK"  # XLSX is a ZIP container, not CSV renamed as .xlsx.
        meta, expert_headers, body, rows = _metadata_and_matrix(data)
        assert meta["VALISTRUCT_PROFILE"] == profile
        assert meta["VALISTRUCT_SCHEMA"] == SCHEMA
        assert len(expert_headers) == judges
        assert expert_headers == [
            f"{'Experto' if profile == 'delphi' else 'Juez'} {i}" for i in range(1, judges + 1)
        ]
        assert len(body) == items
        assert any(row and row[0] == "IDENTIFICACIÓN DE EXPERTOS" for row in rows)


def test_real_roundtrip_keeps_method_specific_responses():
    cvi, _, _ = _make_book("cvi")
    _, _, cvi_rows, _ = _metadata_and_matrix(cvi)
    assert list(cvi_rows[0][3:7]) == [4, 4, 3, 4]

    lawshe, _, _ = _make_book("lawshe")
    _, _, lawshe_rows, _ = _metadata_and_matrix(lawshe)
    assert list(lawshe_rows[1][3:7]) == ["Esencial", "Útil pero no esencial", "Esencial", "No necesario"]

    delphi, _, _ = _make_book("delphi")
    meta, _, delphi_rows, _ = _metadata_and_matrix(delphi)
    assert meta["Modalidad"] == "Delphi modificado"
    assert meta["Ronda"] == 1
    assert list(delphi_rows[0][4:8]) == [5, 4, 4, 5]
